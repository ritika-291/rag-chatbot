import os
import backend.globals as g

from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings
from langchain_community.vectorstores import FAISS
from dotenv import load_dotenv

load_dotenv()
from backend.config import Settings
settings = Settings()

_embeddings_client = None


def get_embeddings():
    global _embeddings_client
    if _embeddings_client is None:
        _embeddings_client = OpenAIEmbeddings(
            model=settings.EMBEDDING_MODEL,
            api_key=settings.LLM_API_KEY,
            base_url=settings.LLM_URL,
            check_embedding_ctx_length=False,
        )
    return _embeddings_client


# ✅ FIX: added session_id param
def _split_and_index(documents, existing_vector=None, original_filename=None, session_id=None):

    computed_chunk_size = settings.CHUNK_SIZE
    computed_chunk_overlap = settings.CHUNK_OVERLAP

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=computed_chunk_size,
        chunk_overlap=computed_chunk_overlap,
    )

    chunks = []
    for doc in documents:
        try:
            doc_chunks = splitter.split_documents([doc])
        except Exception:
            doc_chunks = []

        orig_text = getattr(doc, "page_content", "") or ""

        for c in doc_chunks:
            try:
                meta = dict(getattr(c, "metadata", {}) or {})
            except Exception:
                meta = {}

            meta_src = (getattr(doc, "metadata", None) or {})
            page_label = meta_src.get("page_label")
            page_num = meta_src.get("page")

            file_name = original_filename or meta_src.get("source_file") or meta_src.get("source")
            if file_name:
                file_name = os.path.basename(file_name)

            paragraph = None
            try:
                chunk_text = (getattr(c, "page_content", "") or "").strip()
                if chunk_text and orig_text:
                    pos = orig_text.find(chunk_text[:200])
                    if pos != -1:
                        paragraph = orig_text[:pos].count("\n\n") + 1
            except Exception:
                pass

            reference = None
            if file_name:
                parts = [f"PDF: {file_name}"]
                if page_label is not None:
                    parts.append(f"Page {page_label}")
                elif page_num is not None:
                    parts.append(f"Page {page_num}")
                if paragraph:
                    parts.append(f"Para {paragraph}")
                reference = "(" + ", ".join(parts) + ")"

            # ✅ FIX: added session_id
            meta.update({
                "source": "pdf",
                "file_name": file_name,
                "page": page_label if page_label is not None else page_num,
                "paragraph": paragraph,
                "reference": reference,
                "session_id": session_id
            })

            c.metadata = meta
            chunks.append(c)

    if not chunks:
        raise ValueError("Document splitting resulted in empty chunks")

    embeddings = get_embeddings()

    vector_db = existing_vector
    if existing_vector is not None:
        try:
            existing_vector.add_documents(chunks)
            vector_db = existing_vector
        except Exception:
            new_vec = FAISS.from_documents(chunks, embeddings)
            try:
                new_vec.merge_from(existing_vector)
                vector_db = new_vec
            except Exception:
                vector_db = new_vec
    else:
        vector_db = FAISS.from_documents(chunks, embeddings)

    return vector_db, chunks


# ✅ FIX: added session_id
def process_pdf(pdf_path: str, existing_vector=None, original_filename=None, session_id=None):
    try:
        loader = PyPDFLoader(pdf_path)

        try:
            pages_iter = loader.lazy_load()
        except Exception:
            pages_iter = iter(loader.load())

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
        )

        embeddings = get_embeddings()

        vector_db = existing_vector
        batch = []
        any_chunks = False
        all_chunks = []
        BATCH_PAGES = settings.BATCH_PAGES

        for page in pages_iter:
            batch.append(page)

            if len(batch) >= BATCH_PAGES:
                page_chunks = []

                for p in batch:
                    try:
                        pcs = splitter.split_documents([p])
                    except Exception:
                        pcs = []

                    orig_text = getattr(p, "page_content", "") or ""

                    for c in pcs:
                        try:
                            meta = dict(getattr(c, "metadata", {}) or {})
                        except Exception:
                            meta = {}

                        meta_src = (getattr(p, "metadata", None) or {})
                        page_label = meta_src.get("page_label")
                        page_num = meta_src.get("page")

                        file_name = original_filename or meta_src.get("source_file") or meta_src.get("source")
                        if file_name:
                            file_name = os.path.basename(file_name)

                        paragraph = None
                        try:
                            chunk_text = (c.page_content or "").strip()
                            if chunk_text and orig_text:
                                pos = orig_text.find(chunk_text[:200])
                                if pos != -1:
                                    paragraph = orig_text[:pos].count("\n\n") + 1
                        except Exception:
                            pass

                        reference = None
                        if file_name:
                            parts = [f"PDF: {file_name}"]
                            if page_label:
                                parts.append(f"Page {page_label}")
                            elif page_num is not None:
                                parts.append(f"Page {page_num}")
                            if paragraph:
                                parts.append(f"Para {paragraph}")
                            reference = "(" + ", ".join(parts) + ")"

                        # ✅ FIX: added session_id
                        meta.update({
                            "source": "pdf",
                            "file_name": file_name,
                            "page": page_label if page_label else page_num,
                            "paragraph": paragraph,
                            "reference": reference,
                            "session_id": session_id
                        })

                        c.metadata = meta
                        page_chunks.append(c)

                batch = []

                if not page_chunks:
                    continue

                all_chunks.extend(page_chunks)
                any_chunks = True

                if vector_db is None:
                    vector_db = FAISS.from_documents(page_chunks, embeddings)
                else:
                    vector_db.add_documents(page_chunks)

        # ✅ flush remaining
        if batch:
            page_chunks = []

            for p in batch:
                pcs = splitter.split_documents([p])
                orig_text = getattr(p, "page_content", "") or ""

                for c in pcs:
                    meta = dict(getattr(c, "metadata", {}) or {})
                    meta_src = (getattr(p, "metadata", None) or {})

                    page_label = meta_src.get("page_label")
                    page_num = meta_src.get("page")

                    file_name = original_filename or meta_src.get("source_file") or meta_src.get("source")
                    if file_name:
                        file_name = os.path.basename(file_name)

                    paragraph = None
                    try:
                        pos = orig_text.find((c.page_content or "")[:200])
                        if pos != -1:
                            paragraph = orig_text[:pos].count("\n\n") + 1
                    except:
                        pass

                    reference = None
                    if file_name:
                        parts = [f"PDF: {file_name}", f"Page {page_label or page_num}"]
                        if paragraph:
                            parts.append(f"Para {paragraph}")
                        reference = "(" + ", ".join(parts) + ")"

                    # ✅ FIX: removed src + added session_id
                    meta.update({
                        "source": "pdf",
                        "file_name": file_name,
                        "page": page_label if page_label else page_num,
                        "paragraph": paragraph,
                        "reference": reference,
                        "session_id": session_id
                    })

                    c.metadata = meta
                    page_chunks.append(c)

            if page_chunks:
                all_chunks.extend(page_chunks)
                any_chunks = True
                if vector_db is None:
                    vector_db = FAISS.from_documents(page_chunks, embeddings)
                else:
                    vector_db.add_documents(page_chunks)

        if not any_chunks:
            raise ValueError("No text extracted from PDF")

        # ✅ FIX: session-aware storage (append to existing chunks if present)
        if session_id is not None:
            existing_chunks = g.session_docs.get(session_id) or []
            g.session_docs[session_id] = existing_chunks + all_chunks
        else:
            existing_chunks = getattr(g, "global_docs", None) or []
            g.global_docs = existing_chunks + all_chunks

        return vector_db, all_chunks

    except Exception as e:
        raise RuntimeError(f"PDF processing failed: {str(e)}")


# ✅ FIX: pass session_id
def process_document(path, existing_vector=None, original_filename=None, session_id=None):
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return process_pdf(
            path,
            existing_vector=existing_vector,
            original_filename=original_filename,
            session_id=session_id
        )

    if ext in (".docx", ".doc"):
        try:
            from docx import Document as DocxDocument
        except Exception as e:
            raise RuntimeError(f"DOCX processing requires python-docx: {e}")

        try:
            from langchain.schema import Document as LCDocument
        except Exception:
            class LCDocument:
                def __init__(self, page_content, metadata=None):
                    self.page_content = page_content
                    self.metadata = metadata or {}

        doc = DocxDocument(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        text = "\n".join(paragraphs)

        if not text.strip():
            raise ValueError("No text extracted from DOCX")

        documents = [LCDocument(page_content=text, metadata={"source_file": os.path.basename(path)})]

        vector, chunks = _split_and_index(
            documents,
            existing_vector=existing_vector,
            original_filename=original_filename,
            session_id=session_id
        )

        try:
            if session_id is not None:
                existing_chunks = g.session_docs.get(session_id) or []
                g.session_docs[session_id] = existing_chunks + chunks
            else:
                existing_chunks = getattr(g, "global_docs", None) or []
                g.global_docs = existing_chunks + chunks
        except Exception:
            pass

        return vector, chunks

    raise RuntimeError("Unsupported file type")